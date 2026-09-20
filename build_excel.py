"""
Genera Busqueda_Empleo.xlsx a partir de las ofertas encontradas en ofertas.json.

Este script es generico: no contiene datos de ningun candidato ni ofertas
hardcodeadas. Los resultados de cada busqueda se guardan en ofertas.json
(ver ofertas.ejemplo.json para el formato esperado) y este script solo se
encarga de darle formato y generar el Excel.
"""

import json
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

BASE_DIR = Path(__file__).parent
DATA_FILE = BASE_DIR / "ofertas.json"
OUTPUT_FILE = BASE_DIR / "Busqueda_Empleo.xlsx"

HEADERS = ["Puesto", "Empresa", "Categoria", "Nivel", "Ubicacion / Modalidad", "Descripcion", "Fuente", "URL"]
FIELDS = ["puesto", "empresa", "categoria", "nivel", "ubicacion", "descripcion", "fuente", "url"]


def cargar_ofertas():
    if not DATA_FILE.exists():
        sys.exit(
            f"No se encontro {DATA_FILE.name}. Genera primero un archivo con las ofertas "
            f"encontradas (ver ofertas.ejemplo.json para el formato) antes de correr este script."
        )
    with open(DATA_FILE, encoding="utf-8") as f:
        payload = json.load(f)
    ofertas = payload.get("ofertas", [])
    if not ofertas:
        sys.exit(f"{DATA_FILE.name} no contiene ofertas. Nada que generar.")
    return payload


def construir_excel(payload):
    ofertas = payload["ofertas"]
    nota = payload.get("nota", "")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ofertas de Empleo"

    header_font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    cell_font = Font(name="Arial", size=10)
    wrap_align = Alignment(wrap_text=True, vertical="top")
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col, title in enumerate(HEADERS, start=1):
        c = ws.cell(row=1, column=col, value=title)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border

    for r, oferta in enumerate(ofertas, start=2):
        for col, field in enumerate(FIELDS, start=1):
            value = oferta.get(field, "")
            cell = ws.cell(row=r, column=col, value=value)
            cell.font = cell_font
            cell.alignment = wrap_align
            cell.border = border
            if field == "url" and value:
                cell.hyperlink = value
                cell.font = Font(name="Arial", size=10, color="0563C1", underline="single")

    widths = {"A": 32, "B": 24, "C": 22, "D": 30, "E": 26, "F": 55, "G": 14, "H": 45}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    ws.row_dimensions[1].height = 24
    for r in range(2, len(ofertas) + 2):
        ws.row_dimensions[r].height = 70

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:H{len(ofertas) + 1}"

    if nota:
        note_row = len(ofertas) + 3
        note = ws.cell(row=note_row, column=1, value=nota)
        note.font = Font(name="Arial", size=9, italic=True, color="808080")
        ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=8)
        note.alignment = Alignment(wrap_text=True)

    wb.save(OUTPUT_FILE)
    print(f"OK: {OUTPUT_FILE.name} generado con {len(ofertas)} ofertas.")


if __name__ == "__main__":
    construir_excel(cargar_ofertas())
